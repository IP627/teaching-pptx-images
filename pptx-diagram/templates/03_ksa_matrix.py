# -*- coding: utf-8 -*-
"""模板03：三维目标矩阵 (KSA — Knowledge/Skill/Attitude Matrix)
来源：国一 交通运输 + 国一 建筑工程 + 国一 视觉传感小车
适用场景：教学目标分解、培养规格矩阵、能力素质模型
结构：顶部分类标签 → 三列知识/技能/素养 → 底部教学重难点
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：三维教学目标体系'
PAGE_SUBTITLE = '知识目标 · 技能目标 · 素养目标  |  对标课程标准与职业技能等级标准'

# 三列目标定义
COLUMNS = [
    {
        'title': '知识目标',
        'icon': '知',
        'color': B1,
        'items': [
            '知识点描述第一行',
            '知识点描述第二行',
            '知识点描述第三行',
            '知识点描述第四行',
        ],
    },
    {
        'title': '技能目标',
        'icon': '能',
        'color': B2,
        'items': [
            '技能点描述第一行',
            '技能点描述第二行',
            '技能点描述第三行',
            '技能点描述第四行',
        ],
    },
    {
        'title': '素养目标',
        'icon': '德',
        'color': B3,
        'items': [
            '素养点描述第一行',
            '素养点描述第二行',
            '素养点描述第三行',
            '素养点描述第四行',
        ],
    },
]

# 教学重难点
TEACHING_FOCUS = [
    '教学重点描述第一行',
    '教学重点描述第二行',
    '教学重点描述第三行',
]

TEACHING_DIFFICULTY = [
    '教学难点描述第一行',
    '教学难点描述第二行',
    '教学难点描述第三行',
]

# 底部能力标注
BOTTOM_CAPABILITY = '核心技术能力：XXXX  |  职业技能等级：XXXX  |  对接赛项：XXXX'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_ksa_matrix():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # 三列
    col_w = 3.85; col_gap = 0.18
    col_start_x = 0.55

    for i, col in enumerate(COLUMNS):
        cx = col_start_x + i*(col_w + col_gap)
        cy_top = 0.95

        # 列标题
        rrect(slide, cx, cy_top, col_w, 0.45, col['color'], radius=0.03)
        # 图标圆圈
        rrect(slide, cx+0.12, cy_top+0.08, 0.3, 0.3, WHITE, radius=0.5)
        tb(slide, cx+0.12, cy_top+0.11, 0.3, 0.24, col['icon'],
           size=11, bold=True, color=col['color'], align=PP_ALIGN.CENTER)
        tb(slide, cx+0.52, cy_top+0.08, col_w-0.7, 0.3, col['title'],
           size=12, bold=True, color=WHITE)

        # 内容卡片
        content_y = cy_top + 0.58
        rrect(slide, cx, content_y, col_w, 3.8, WHITE, border=LINE)
        rect(slide, cx, content_y+0.08, col_w, 0.04, col['color'])

        for j, item in enumerate(col['items']):
            iy = content_y + 0.3 + j*0.85
            # 序号
            rrect(slide, cx+0.15, iy, 0.28, 0.28, col['color'], radius=0.5)
            tb(slide, cx+0.15, iy+0.02, 0.28, 0.24, str(j+1),
               size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            # 描述
            rrect(slide, cx+0.55, iy-0.02, col_w-0.8, 0.7, CARD_LT, border=None, radius=0.03)
            mtb(slide, cx+0.65, iy+0.05, col_w-1.0, 0.6,
                [(item, 8, False, MED, PP_ALIGN.LEFT)])

    # ---- 教学重难点 ----
    focus_y = 5.55
    # 教学重点
    rrect(slide, 0.55, focus_y, 6.0, 1.3, WHITE, border=LINE)
    rect(slide, 0.55, focus_y, 0.07, 1.3, DARK_BLUE)
    tb(slide, 0.8, focus_y+0.1, 2.0, 0.22, '教学重点', size=10, bold=True, color=DARK_BLUE)
    for j, item in enumerate(TEACHING_FOCUS):
        tb(slide, 1.0, focus_y+0.4+j*0.25, 5.4, 0.2, f'· {item}', size=8, color=DARK)

    # 教学难点
    rrect(slide, 6.75, focus_y, 6.0, 1.3, WHITE, border=LINE)
    rect(slide, 6.75, focus_y, 0.07, 1.3, ORANGE)
    tb(slide, 7.0, focus_y+0.1, 2.0, 0.22, '教学难点', size=10, bold=True, color=ORANGE)
    for j, item in enumerate(TEACHING_DIFFICULTY):
        tb(slide, 7.2, focus_y+0.4+j*0.25, 5.4, 0.2, f'· {item}', size=8, color=DARK)

    # 底部能力栏
    cap_y = 7.05
    rect(slide, 0.55, cap_y, 12.25, 0.26, CARD_LT)
    tb(slide, 0.65, cap_y+0.04, 12.0, 0.18, BOTTOM_CAPABILITY,
       size=7.5, color=DARK_BLUE, align=PP_ALIGN.CENTER)

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_ksa_matrix()
    prs.save(os.path.join(out_dir, '模板03-三维目标矩阵.pptx'))
    print('已生成: 模板03-三维目标矩阵.pptx')
