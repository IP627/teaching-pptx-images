# -*- coding: utf-8 -*-
"""模板14：对比映照双栏 (Comparison Dual Column)
来源：200页模板 Slide 45/165 + 明大德 Slide 4/13 + 建筑工程 Slide 3
适用场景：教学重点vs难点、课前vs课后对比、传统vs创新对比、学情前后测
结构：左右双色对比 → 中间差异标注 → 底部综合结论
视觉增强：双色系区分 + 阴影卡片 + 对比连接线 + 差异标注
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
CARD_WARM = RGBColor(0xFD, 0xF0, 0xE8)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)
GREEN     = RGBColor(0x60, 0xBC, 0x90)
RED_ACCENT = RGBColor(0xD9, 0x4F, 0x4F)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def _get_spPr(shape):
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    spPr = shape._element.find(f'{ns}spPr')
    if spPr is None:
        spPr = etree.SubElement(shape._element, f'{ns}spPr')
    return spPr

def add_shadow(shape, blur_rad=35000, dist=18000, alpha=18000):
    spPr = _get_spPr(shape)
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    effectLst = spPr.find(f'{{{ns_a}}}effectLst')
    if effectLst is None:
        effectLst = etree.SubElement(spPr, f'{{{ns_a}}}effectLst')
    outerShdw = etree.SubElement(effectLst, f'{{{ns_a}}}outerShdw', {
        'blurRad': str(blur_rad), 'dist': str(dist),
        'dir': '5400000', 'algn': 'ctr',
    })
    srgbClr = etree.SubElement(outerShdw, f'{{{ns_a}}}srgbClr', {'val': '000000'})
    etree.SubElement(srgbClr, f'{{{ns_a}}}alpha', {'val': str(alpha)})

def set_gradient_fill(shape, c1_hex, c2_hex, angle=90):
    spPr = _get_spPr(shape)
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    for sf in spPr.findall(f'{{{ns_a}}}solidFill'):
        spPr.remove(sf)
    gradFill = etree.SubElement(spPr, f'{{{ns_a}}}gradFill', {'rotWithShape': '1'})
    gsLst = etree.SubElement(gradFill, f'{{{ns_a}}}gsLst')
    for pos, clr in [('0', c1_hex), ('100000', c2_hex)]:
        gs = etree.SubElement(gsLst, f'{{{ns_a}}}gs', {'pos': pos})
        etree.SubElement(gs, f'{{{ns_a}}}srgbClr', {'val': clr})
    lin = etree.SubElement(gradFill, f'{{{ns_a}}}lin', {'ang': str(angle * 60000), 'scaled': '1'})

def rrect(s, l, t, w, h, fill, border=None, radius=0.04, shadow=False, gradient=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    if gradient:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        set_gradient_fill(sh, gradient[0], gradient[1])
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    if shadow: add_shadow(sh)
    return sh

def rect(s, l, t, w, h, fill, border=None, shadow=False):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    if shadow: add_shadow(sh)
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：教学重难点对比分析'
PAGE_SUBTITLE = '左：教学重点  |  右：教学难点  |  中间：突破策略'

# 左侧列
LEFT_COLUMN = {
    'title': '教学重点',
    'sub': 'Teaching Focus',
    'color': B1,
    'accent_color': B2,
    'bg_tint': CARD_LT,
    'items': [
        'XX基本概念与核心原理',
        'XX工艺流程及关键技术要点',
        'XX系统的设计方法与规范',
        'XX与YY的协同工作机制',
    ],
}

# 右侧列
RIGHT_COLUMN = {
    'title': '教学难点',
    'sub': 'Teaching Difficulty',
    'color': ORANGE,
    'accent_color': RED_ACCENT,
    'bg_tint': CARD_WARM,
    'items': [
        '复杂工况下的XX故障诊断',
        'XX算法的数学原理与推导',
        '多因素耦合的XX优化问题',
        '实际工程场景的灵活应用',
    ],
}

# 中间突破策略
STRATEGIES = [
    {'icon': '虚', 'text': '虚拟仿真\n平台演练', 'color': B3},
    {'icon': '案', 'text': '典型案例\n拆解分析', 'color': B4},
    {'icon': '导', 'text': '双师协同\n分层指导', 'color': GREEN},
    {'icon': '评', 'text': '过程评价\n即时反馈', 'color': B2},
]

BOTTOM_SLOGAN = '"聚焦重点 · 攻克难点"'
BOTTOM_SUB = '基于学情诊断的精准教学策略'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_comparison():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 渐变标题栏
    rrect(slide, 0.65, 0.18, 12.0, 0.48, DARK_BLUE, radius=0.04,
          gradient=('1E5793', '2E78B5'))
    tb(slide, 0.85, 0.22, 8.0, 0.28, PAGE_TITLE, size=16, bold=True, color=WHITE)
    tb(slide, 0.85, 0.48, 10.0, 0.16, PAGE_SUBTITLE, size=7.5, color=RGBColor(0xB0,0xCC,0xE5))

    # ---- 两列对比布局 ----
    col_w = 4.35; col_gap = 1.0
    col_top = 0.85
    col_h = 4.1
    left_x = 0.55
    mid_x = left_x + col_w + 0.15
    right_x = mid_x + col_gap + 0.15

    # 左列背景
    rrect(slide, left_x, col_top, col_w, col_h, WHITE, border=LINE, shadow=True)
    # 左列标题
    rrect(slide, left_x, col_top, col_w, 0.52, LEFT_COLUMN['color'], radius=0.03,
          gradient=('1E5793', '2E78B5'))
    tb(slide, left_x+0.12, col_top+0.05, col_w-0.24, 0.22, LEFT_COLUMN['title'],
       size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, left_x+0.12, col_top+0.3, col_w-0.24, 0.16, LEFT_COLUMN['sub'],
       size=7, color=RGBColor(0xB0,0xCC,0xE5), align=PP_ALIGN.CENTER)

    # 左列内容条目
    for i, item in enumerate(LEFT_COLUMN['items']):
        iy = col_top + 0.7 + i*0.78
        rrect(slide, left_x+0.15, iy, col_w-0.3, 0.64, LEFT_COLUMN['bg_tint'], border=None, radius=0.03)
        rrect(slide, left_x+0.22, iy+0.15, 0.32, 0.32, LEFT_COLUMN['color'], radius=0.5)
        tb(slide, left_x+0.22, iy+0.19, 0.32, 0.24, str(i+1),
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left_x+0.65, iy+0.15, col_w-0.9, 0.34, item, size=9, color=DARK)

    # 右列背景
    rrect(slide, right_x, col_top, col_w, col_h, WHITE, border=LINE, shadow=True)
    # 右列标题
    rrect(slide, right_x, col_top, col_w, 0.52, RIGHT_COLUMN['color'], radius=0.03,
          gradient=('E07030', 'F08850'))
    tb(slide, right_x+0.12, col_top+0.05, col_w-0.24, 0.22, RIGHT_COLUMN['title'],
       size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, right_x+0.12, col_top+0.3, col_w-0.24, 0.16, RIGHT_COLUMN['sub'],
       size=7, color=RGBColor(0xFF,0xE0,0xCC), align=PP_ALIGN.CENTER)

    # 右列内容条目
    for i, item in enumerate(RIGHT_COLUMN['items']):
        iy = col_top + 0.7 + i*0.78
        rrect(slide, right_x+0.15, iy, col_w-0.3, 0.64, RIGHT_COLUMN['bg_tint'], border=None, radius=0.03)
        rrect(slide, right_x+0.22, iy+0.15, 0.32, 0.32, RIGHT_COLUMN['color'], radius=0.5)
        tb(slide, right_x+0.22, iy+0.19, 0.32, 0.24, str(i+1),
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, right_x+0.65, iy+0.15, col_w-0.9, 0.34, item, size=9, color=DARK)

    # ---- 中间：突破策略 ----
    mid_w = col_gap + 0.1

    # VS 分隔
    rrect(slide, mid_x+0.05, col_top+1.8, 0.9, 0.9, WHITE, border=LINE, shadow=True, radius=0.5)
    mtb(slide, mid_x+0.08, col_top+2.02, 0.84, 0.46, [
        ('VS', 14, True, ORANGE, PP_ALIGN.CENTER),
    ])

    # 中间连接线（左侧 → VS → 右侧）
    for i, (ly, ry) in enumerate([(col_top+1.0, col_top+1.0), (col_top+2.6, col_top+2.6)]):
        # 小箭头/连线
        rect(slide, left_x+col_w+0.02, ly, 0.15, 0.01, LINE)
        rect(slide, right_x-0.17, ry, 0.15, 0.01, LINE)

    # 策略标签
    strategy_top = col_top + 3.1
    rect(slide, mid_x-0.02, strategy_top-0.8, mid_w+0.04, 0.01, LINE)
    tb(slide, mid_x-0.1, strategy_top-0.68, mid_w+0.2, 0.18, '突破策略',
       size=8, bold=True, color=MED_BLUE, align=PP_ALIGN.CENTER)

    for i, st in enumerate(STRATEGIES):
        sy = strategy_top + i*0.28
        sx = mid_x - 0.05
        rrect(slide, sx+0.02, sy+0.01, 0.22, 0.22, st['color'], radius=0.5)
        tb(slide, sx+0.02, sy+0.03, 0.22, 0.18, st['icon'],
           size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, sx+0.3, sy, mid_w-0.2, 0.24, st['text'],
           size=6.5, color=DARK, align=PP_ALIGN.CENTER)

    # ---- 底部结论 ----
    conclusion_y = col_top + col_h + 0.12
    rrect(slide, 0.55, conclusion_y, 12.15, 0.65, WHITE, border=LINE, shadow=True)
    rect(slide, 0.55, conclusion_y, 12.15, 0.04, ORANGE)
    mtb(slide, 0.7, conclusion_y+0.08, 11.8, 0.5, [
        (BOTTOM_SLOGAN, 13, True, ORANGE, PP_ALIGN.CENTER),
        (BOTTOM_SUB, 8, False, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_comparison()
    prs.save(os.path.join(out_dir, '模板14-对比映照双栏.pptx'))
    print('已生成: 模板14-对比映照双栏.pptx')
