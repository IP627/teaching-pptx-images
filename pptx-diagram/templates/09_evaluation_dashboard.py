# -*- coding: utf-8 -*-
"""模板09：评价体系仪表盘 (Evaluation Dashboard)
来源：劳动教育 Slide 5 + 国二 思政-明大德 Slide 11 + 国一 建筑工程 Slide 9 + 国一 视觉传感小车 Slide 10
适用场景：多维教学评价体系、考核方案设计、评价主体与权重分布
结构：评价维度分区 → 权重进度条 → 评价主体 → 评价手段
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)
GREEN     = RGBColor(0x60, 0xBC, 0x90)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：多维教学评价体系'
PAGE_SUBTITLE = '过程评价 · 结果评价 · 增值评价 · 综合评价  |  多元主体 · 多维指标'

# 评价大类
EVALUATION_CATEGORIES = [
    {
        'title': '过程评价',
        'weight': '60%',
        'color': B1,
        'sub_items': [
            {'name': '课前学习数据', 'weight': '10%', 'desc': '平台学习时长 · 预练完成度 · 测验成绩'},
            {'name': '课堂表现', 'weight': '15%', 'desc': '小组协作 · 实操表现 · 课堂互动'},
            {'name': '项目过程', 'weight': '20%', 'desc': '方案设计 · 阶段性成果 · 过程记录'},
            {'name': '素养表现', 'weight': '15%', 'desc': '工匠精神 · 团队意识 · 创新能力'},
        ],
    },
    {
        'title': '结果评价',
        'weight': '40%',
        'color': B2,
        'sub_items': [
            {'name': '知识考核', 'weight': '10%', 'desc': '理论知识 · 期末测试'},
            {'name': '技能考核', 'weight': '15%', 'desc': '实操考核 · 项目成品质量'},
            {'name': '证书/大赛', 'weight': '15%', 'desc': '1+X证书 · 技能大赛成绩'},
        ],
    },
]

# 评价主体
EVALUATORS = [
    {'name': '教师评价', 'icon': '师', 'color': B1},
    {'name': '学生自评', 'icon': '自', 'color': B2},
    {'name': '小组互评', 'icon': '互', 'color': B3},
    {'name': '企业导师', 'icon': '企', 'color': B4},
    {'name': '平台数据', 'icon': '数', 'color': GREEN},
]

# 评价手段
EVAL_METHODS = ['考核量表', '学习画像', '标准分差值', '成长曲线', '雷达图分析']

# 底部标语
BOTTOM_SLOGAN = '"以评促学 · 以评促教"'
BOTTOM_SUB = '四评五测 · 全过程数据驱动的多元评价体系'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_evaluation_dashboard():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # ---- 左半区：过程评价 ----
    left_x, left_w = 0.55, 6.0
    right_x, right_w = 6.8, 6.0
    item_top = 0.92

    for ci, cat in enumerate(EVALUATION_CATEGORIES):
        cx = left_x if ci == 0 else right_x
        cw = left_w if ci == 0 else right_w
        cc = cat['color']

        # 类别标题
        rrect(slide, cx, item_top, cw, 0.42, cc, radius=0.03)
        tb(slide, cx+0.12, item_top+0.05, cw-1.5, 0.24, cat['title'],
           size=12, bold=True, color=WHITE)
        # 权重标签
        rrect(slide, cx+cw-1.1, item_top+0.07, 0.9, 0.28, WHITE, radius=0.03)
        tb(slide, cx+cw-1.1, item_top+0.09, 0.9, 0.24, f'占比 {cat["weight"]}',
           size=9, bold=True, color=cc, align=PP_ALIGN.CENTER)

        # 子项
        sub_top = item_top + 0.55
        sub_h = 0.72; sub_gap = 0.12

        for si, sub in enumerate(cat['sub_items']):
            sy = sub_top + si*(sub_h + sub_gap)
            # 子项背景
            rrect(slide, cx, sy, cw, sub_h, WHITE, border=LINE)
            # 左侧色条
            rect(slide, cx, sy+0.06, 0.05, sub_h-0.12, cc)
            # 名称
            tb(slide, cx+0.15, sy+0.05, 1.8, 0.2, sub['name'],
               size=9, bold=True, color=DARK)
            # 子权重
            rrect(slide, cx+cw-1.1, sy+0.06, 0.9, 0.22, CARD_LT, radius=0.03)
            tb(slide, cx+cw-1.1, sy+0.07, 0.9, 0.2, sub['weight'],
               size=8, bold=True, color=MED_BLUE, align=PP_ALIGN.CENTER)
            # 描述
            tb(slide, cx+0.15, sy+0.3, cw-1.4, 0.18, sub['desc'],
               size=7, color=MED)
            # 权重进度条
            pct = int(sub['weight'].replace('%', ''))
            bar_w = (cw-1.6)*(pct/100)
            rect(slide, cx+0.15, sy+sub_h-0.12, bar_w, 0.04, cc)

    # ---- 底部左：评价主体 ----
    evaluator_y = item_top + 0.55 + 4*(0.72+0.12) + 0.2
    tb(slide, 0.65, evaluator_y, 1.5, 0.2, '评价主体：', size=8, bold=True, color=MED_BLUE)
    etop = evaluator_y + 0.25

    for i, ev in enumerate(EVALUATORS):
        ex = 0.65 + i*2.4
        # 主体卡片
        rrect(slide, ex, etop, 2.12, 0.6, WHITE, border=LINE)
        rrect(slide, ex+0.08, etop+0.12, 0.36, 0.36, ev['color'], radius=0.5)
        tb(slide, ex+0.08, etop+0.16, 0.36, 0.28, ev['icon'],
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, ex+0.52, etop+0.14, 1.45, 0.22, ev['name'],
           size=9, bold=True, color=DARK)

    # ---- 底部右：评价手段 ----
    method_y = evaluator_y
    tb(slide, 6.9, method_y, 1.5, 0.2, '评价手段：', size=8, bold=True, color=MED_BLUE)

    for i, mn in enumerate(EVAL_METHODS):
        mx = 6.9 + i*1.1
        mw = 0.95
        rrect(slide, mx, etop, mw, 0.34, CARD_LT, border=None, radius=0.03)
        tb(slide, mx+0.04, etop+0.05, mw-0.08, 0.24, mn,
           size=7, color=MED_BLUE, align=PP_ALIGN.CENTER)

    # ---- 评价体系总览架构图 ----
    arch_y = etop + 0.85
    # 中心环
    center_x = 6.0
    rrect(slide, center_x, arch_y, 1.1, 0.5, DARK_BLUE, radius=0.06)
    tb(slide, center_x+0.05, arch_y+0.06, 1.0, 0.38, '综合评价',
       size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 四个环绕模块
    arch_items = [
        ('过程评价', 0.55, B1),
        ('结果评价', 4.0, B2),
        ('增值评价', 7.5, B3),
        ('综合素养', 10.0, B4),
    ]
    for label, ax, ac in arch_items:
        rrect(slide, ax, arch_y-0.22, 1.6, 0.44, ac, radius=0.04)
        tb(slide, ax, arch_y-0.14, 1.6, 0.28, label,
           size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # 连接箭头
        if ax < center_x:
            arrow_r(slide, ax+1.62, arch_y-0.02, 0.18, 0.1, ac)
        else:
            arrow_r(slide, ax-0.25, arch_y-0.02, 0.18, 0.1, ac)

    # ---- 底部标语 ----
    slogan_y = arch_y + 0.65
    rrect(slide, 3.0, slogan_y, 7.3, 0.6, WHITE, border=LINE)
    rect(slide, 3.0, slogan_y, 7.3, 0.04, ORANGE)
    mtb(slide, 3.15, slogan_y+0.06, 7.0, 0.48, [
        (BOTTOM_SLOGAN, 12, True, ORANGE, PP_ALIGN.CENTER),
        (BOTTOM_SUB, 8, False, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_evaluation_dashboard()
    prs.save(os.path.join(out_dir, '模板09-评价体系仪表盘.pptx'))
    print('已生成: 模板09-评价体系仪表盘.pptx')
