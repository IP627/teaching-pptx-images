# -*- coding: utf-8 -*-
"""模板06：教学流程时序 (Teaching Process Timeline)
来源：Z62 10步流水线 + 国一 视觉传感小车 Slide 5 + 国二 思政-明大德 Slide 5 + 劳动教育 Slide 7/9
适用场景：课前/课中/课后三段式教学流程、教学实施步骤、闯关式教学
结构：上排阶段标签 → 流程步骤节点 → 底部教学策略/方法标注
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

def arrow_d(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：教学实施流程'
PAGE_SUBTITLE = '课前探学 → 课中做学 → 课后拓学  |  课程名称 · X学时'

# 三大阶段
PHASES = [
    {
        'name': '课前探学',
        'tag': '线上',
        'color': B2,
        'steps': [
            {'label': '任务发布', 'desc': '平台推送学习\n资源与任务单'},
            {'label': '自主预练', 'desc': '学生分组完成\n预练任务'},
            {'label': '学情分析', 'desc': '教师分析预练\n数据调整教学'},
        ],
    },
    {
        'name': '课中做学',
        'tag': '线下+线上',
        'color': B1,
        'steps': [
            {'label': '情境导入', 'desc': '企业真实案例\n引出任务'},
            {'label': '任务驱动', 'desc': '分组协作探究\n实操演练'},
            {'label': '成果展示', 'desc': '小组汇报互评\n教师点评'},
            {'label': '总结提升', 'desc': '重难点强化\n知识内化'},
        ],
    },
    {
        'name': '课后拓学',
        'tag': '线上+实践',
        'color': B3,
        'steps': [
            {'label': '拓展任务', 'desc': '企业真实项目\n拓展练习'},
            {'label': '反思总结', 'desc': '学习报告撰写\n知识复盘'},
        ],
    },
]

# 底部教学策略
STRATEGIES = [
    {'name': '任务驱动', 'icon': '驱'},
    {'name': '案例教学', 'icon': '案'},
    {'name': '小组协作', 'icon': '协'},
    {'name': '虚实结合', 'icon': '虚'},
    {'name': '双师导学', 'icon': '导'},
]

# 底部标语
BOTTOM_SLOGAN = '"核心标语"'
BOTTOM_SUB = '以学生为中心 · 以产出为导向'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_teaching_timeline():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # ---- 三阶段布局 ----
    # 计算每个阶段的水平空间
    total_w = 11.8
    start_x = 0.9
    phase_gap = 0.2

    # 各阶段宽度按步骤数比例分配
    total_steps = sum(len(p['steps']) for p in PHASES)
    phase_widths = [len(p['steps'])/total_steps * total_w for p in PHASES]

    px = start_x
    phase_top = 0.92

    for pi, phase in enumerate(PHASES):
        pw = phase_widths[pi]
        ac = phase['color']

        # 阶段背景
        rrect(slide, px, phase_top, pw, 4.3, WHITE, border=LINE, radius=0.04)
        # 阶段标题栏
        rrect(slide, px, phase_top, pw, 0.4, ac, radius=0.03)
        tb(slide, px+0.08, phase_top+0.05, pw-0.6, 0.22, phase['name'],
           size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, px+0.08, phase_top+0.24, pw-0.6, 0.14, phase['tag'],
           size=7, color=RGBColor(0xB0,0xCC,0xE5), align=PP_ALIGN.CENTER)

        # 步骤节点
        step_top = phase_top + 0.55
        step_w = pw - 0.26
        step_h = 1.0
        step_gap = 0.15

        for si, step in enumerate(phase['steps']):
            sy = step_top + si*(step_h + step_gap)
            # 步骤卡片
            rrect(slide, px+0.13, sy, step_w, step_h, CARD_LT, border=None, radius=0.03)
            # 步骤序号圆
            rrect(slide, px+0.22, sy+0.12, 0.34, 0.34, ac, radius=0.5)
            tb(slide, px+0.22, sy+0.16, 0.34, 0.26, str(si+1),
               size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            # 步骤标签与描述
            tb(slide, px+0.7, sy+0.08, step_w-0.9, 0.2, step['label'],
               size=9, bold=True, color=DARK)
            mtb(slide, px+0.7, sy+0.32, step_w-0.9, step_h-0.4,
                [(step['desc'], 7, False, MED, PP_ALIGN.LEFT)])

            # 步骤间箭头
            if si < len(phase['steps'])-1:
                arrow_d(slide, px+pw/2-0.06, sy+step_h+0.01, 0.12, 0.08, ac)

        px += pw + phase_gap

    # ---- 底部教学策略标签 ----
    strategy_y = 5.5
    rect(slide, 0.55, strategy_y, 12.25, 0.01, LINE)

    # 教学策略标题
    tb(slide, 0.65, strategy_y+0.08, 1.2, 0.2, '教学策略：', size=8, bold=True, color=MED_BLUE)

    sx_start = 1.85
    for i, st in enumerate(STRATEGIES):
        sx = sx_start + i*1.9
        rrect(slide, sx, strategy_y+0.06, 0.28, 0.28, B1, radius=0.5)
        tb(slide, sx, strategy_y+0.08, 0.28, 0.22, st['icon'],
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, sx+0.34, strategy_y+0.08, 1.5, 0.22, st['name'],
           size=8, color=DARK)

    # ---- 底部标语栏 ----
    slogan_y = 5.95
    rrect(slide, 3.0, slogan_y, 7.3, 0.7, WHITE, border=LINE)
    rect(slide, 3.0, slogan_y, 7.3, 0.04, ORANGE)
    mtb(slide, 3.15, slogan_y+0.10, 7.0, 0.5, [
        (BOTTOM_SLOGAN, 13, True, ORANGE, PP_ALIGN.CENTER),
        (BOTTOM_SUB, 9, False, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    # 底部时间轴横线
    timeline_y = 6.85
    rect(slide, 0.9, timeline_y, 11.5, 0.012, MED_BLUE)
    for i in range(0, 12):
        dot_x = 0.9 + i*1.05
        rrect(slide, dot_x, timeline_y-0.04, 0.08, 0.08, MED_BLUE, radius=0.5)

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_teaching_timeline()
    prs.save(os.path.join(out_dir, '模板06-教学流程时序.pptx'))
    print('已生成: 模板06-教学流程时序.pptx')
