# -*- coding: utf-8 -*-
"""模板02：四阶递进流 (Four-Stage Progressive Flow)
来源：国二 思政-明大德 + 国一 交通运输 + 国一 建筑工程
适用场景：四阶递进课程体系、能力进阶路径、内容重构逻辑
结构：入口 → 阶段1 → 阶段2 → 阶段3 → 阶段4 → 出口 (纵向递进)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

# ---- 配色 ----
BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

# ---- 工具函数 ----
def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_d(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X："四阶递进"课程体系'
PAGE_SUBTITLE = '产业认知 → 原理认知 → 模型训练 → 边缘部署 → 场景实战 → 综合融通'

# 左侧能力/需求列表
LEFT_LABEL = '岗位核心能力'
LEFT_ITEMS = [
    {'name': '能力一', 'desc': '能力描述文字', 'color': B1},
    {'name': '能力二', 'desc': '能力描述文字', 'color': B2},
    {'name': '能力三', 'desc': '能力描述文字', 'color': B3},
    {'name': '能力四', 'desc': '能力描述文字', 'color': B4},
]

# 四阶递进阶段
STAGES = [
    {'label': '第1阶', 'name': '原理认知', 'module': '模块X：模块名称', 'hours': 'X学时',
     'desc': 'YOLO演进 · 网络结构 · Anchor机制 · 数据标注', 'color': B1},
    {'label': '第2阶', 'name': '模型训练', 'module': '模块X：模块名称', 'hours': 'X学时',
     'desc': '数据集构建 · 模型训练 · 评估优化 · 调参策略', 'color': B2},
    {'label': '第3阶', 'name': '边缘部署', 'module': '模块X：模块名称', 'hours': 'X学时',
     'desc': 'PyTorch→ONNX→TensorRT · FP32/16/INT8量化', 'color': B3},
    {'label': '第4阶', 'name': '场景实战', 'module': '模块X：模块名称', 'hours': 'X学时',
     'desc': '企业现场教学 · 四类AI算法 · 全流程实战', 'color': B4},
]

# 右侧产出
RIGHT_TITLE = '培养目标'
RIGHT_SLOGAN = '"核心标语"'
RIGHT_SUB = '复合型XXX人才'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_four_stage():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # 三区标题
    HEADER_Y = 0.9
    for x, w, title, sub in [
        (0.55, 2.8, '岗位能力驱动', '产业需求 → 核心能力'),
        (3.55, 6.1, '四阶递进体系', '入口→阶段→出口'),
        (9.85, 2.9, '培养目标产出', '岗课赛证融通'),
    ]:
        rect(slide, x, HEADER_Y, w, 0.32, DARK_BLUE)
        tb(slide, x, HEADER_Y+0.02, w, 0.17, title, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, x, HEADER_Y+0.18, w, 0.13, sub, size=7, color=RGBColor(0xB0,0xCC,0xE5), align=PP_ALIGN.CENTER)

    CONTENT_TOP = HEADER_Y + 0.5

    # ---- 左列：能力卡片 ----
    lx, lw = 0.55, 2.8
    for i, item in enumerate(LEFT_ITEMS):
        cy = CONTENT_TOP + i * 1.22
        ac = item['color']
        rrect(slide, lx, cy, lw, 1.08, WHITE, border=LINE)
        rect(slide, lx+0.06, cy+0.08, 0.05, 0.92, ac)
        rrect(slide, lx+0.22, cy+0.2, 0.42, 0.42, ac, radius=0.5)
        tb(slide, lx+0.22, cy+0.26, 0.42, 0.3, f'{i+1:02d}',
           size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, lx+0.78, cy+0.15, lw-1.0, 0.22, item['name'],
           size=10, bold=True, color=DARK)
        tb(slide, lx+0.78, cy+0.45, lw-1.0, 0.3, item['desc'],
           size=7.5, color=MED)
        # 汇聚箭头
        arrow_r(slide, lx+lw+0.02, cy+0.4, 0.14, 0.12, ac)

    # ---- 中列：四阶递进 ----
    mx, mw = 3.55, 6.1
    # 入口
    entry_y = CONTENT_TOP
    rrect(slide, mx+0.4, entry_y, mw-0.8, 0.42, CARD_LT, border=None, radius=0.03)
    tb(slide, mx+0.5, entry_y+0.06, 0.4, 0.2, '入口', size=7, color=MED_BLUE)
    tb(slide, mx+1.0, entry_y+0.06, 2.0, 0.2, '模块一：起始模块', size=8.5, bold=True, color=DARK)
    arrow_d(slide, mx+mw/2-0.07, entry_y+0.48, 0.14, 0.12, MED_BLUE)

    # 四阶段
    stage_top = entry_y + 0.66
    stage_h = 0.95; stage_gap = 0.18

    for i, st in enumerate(STAGES):
        sy = stage_top + i*(stage_h+stage_gap)
        sc = st['color']
        rrect(slide, mx, sy, mw, stage_h, WHITE, border=LINE)
        rect(slide, mx, sy+0.06, 0.07, stage_h-0.12, sc)
        rrect(slide, mx+0.16, sy+0.12, 0.62, 0.3, sc, radius=0.03)
        tb(slide, mx+0.16, sy+0.14, 0.62, 0.26, st['label'],
           size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, mx+0.95, sy+0.08, 1.3, 0.22, st['name'],
           size=12, bold=True, color=DARK)
        tb(slide, mx+2.3, sy+0.12, 3.0, 0.18, st['module'],
           size=9.5, bold=True, color=sc)
        rrect(slide, mx+mw-1.2, sy+0.1, 0.95, 0.28, CARD_LT, radius=0.03)
        tb(slide, mx+mw-1.2, sy+0.12, 0.95, 0.24, st['hours'],
           size=8, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
        # 关键词标签
        keywords = st['desc'].split(' · ')
        kx = mx + 0.95
        for kw in keywords:
            kw_w = len(kw)*0.11 + 0.18
            rrect(slide, kx, sy+stage_h-0.32, kw_w, 0.22, CARD_LT, radius=0.02)
            tb(slide, kx+0.04, sy+stage_h-0.31, kw_w-0.08, 0.2, kw, size=6.5, color=MED_BLUE)
            kx += kw_w + 0.06
        if i < len(STAGES)-1:
            arrow_d(slide, mx+0.18, sy+stage_h+0.02, 0.12, 0.11, sc)

    # 出口
    exit_y = stage_top + len(STAGES)*(stage_h+stage_gap) - stage_gap + 0.06
    rrect(slide, mx+0.4, exit_y, mw-0.8, 0.42, DARK_BLUE, border=None, radius=0.03)
    tb(slide, mx+0.5, exit_y+0.06, 0.4, 0.2, '出口', size=7, color=RGBColor(0xB0,0xCC,0xE5))
    tb(slide, mx+1.0, exit_y+0.06, 2.5, 0.2, '模块N：综合融通', size=8.5, bold=True, color=WHITE)
    tb(slide, mx+mw-1.3, exit_y+0.06, 1.1, 0.2, 'X学时', size=7.5, color=RGBColor(0xB0,0xCC,0xE5))

    # 递进标语
    tb(slide, mx, exit_y+0.5, mw, 0.2,
       f'▲ 四阶递进：{PAGE_SUBTITLE} ▲', size=7.5, color=MED_BLUE, align=PP_ALIGN.CENTER)

    # ---- 右列：培养目标 ----
    rx, rw = 9.85, 2.9
    goal_y = CONTENT_TOP
    goal_h = 1.8
    rrect(slide, rx, goal_y, rw, goal_h, WHITE, border=LINE)
    rect(slide, rx, goal_y, rw, 0.05, ORANGE)
    mtb(slide, rx+0.12, goal_y+0.15, rw-0.24, goal_h-0.3, [
        (RIGHT_TITLE, 9, False, MED_BLUE, PP_ALIGN.CENTER),
        ('', 5, False, DARK, PP_ALIGN.CENTER),
        (RIGHT_SLOGAN, 14, True, ORANGE, PP_ALIGN.CENTER),
        ('', 5, False, DARK, PP_ALIGN.CENTER),
        (RIGHT_SUB, 11, True, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    # 产出能力卡片
    out_top = goal_y + goal_h + 0.15
    for j, (otitle, ocolor) in enumerate([('知识目标', B1), ('技能目标', B2), ('素质目标', B3)]):
        oy = out_top + j*0.95
        rrect(slide, rx, oy, rw, 0.8, WHITE, border=LINE)
        rect(slide, rx+0.08, oy+0.1, 0.05, 0.6, ocolor)
        tb(slide, rx+0.22, oy+0.1, rw-0.35, 0.2, otitle, size=9, bold=True, color=ocolor)
        tb(slide, rx+0.22, oy+0.38, rw-0.35, 0.3, '目标描述第一行\n目标描述第二行', size=7, color=MED)

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_four_stage()
    prs.save(os.path.join(out_dir, '模板02-四阶递进流.pptx'))
    print('已生成: 模板02-四阶递进流.pptx')
