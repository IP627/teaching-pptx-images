# -*- coding: utf-8 -*-
"""模板17：封面目录导航页 (Cover & Navigation Page)
来源：200页模板 Slide 1-2 + 29优秀模板 Slide 1 + 重构框架 Slide 1
适用场景：实施报告封面、章节标题页、目录导航、分隔过渡页
结构：大标题区域 + 装饰分割线 + 项目信息 + 目录列表（可选）
视觉增强：渐变背景区 + 装饰角标 + 阴影分隔线
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)
GREEN     = RGBColor(0x60, 0xBC, 0x90)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def _get_spPr(shape):
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    spPr = shape._element.find(f'{ns}spPr')
    if spPr is None:
        spPr = etree.SubElement(shape._element, f'{ns}spPr')
    return spPr

def add_shadow(shape, blur_rad=35000, dist=18000, alpha=18000):
    spPr = _get_spPr(shape)
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    effectLst = spPr.find(f'{{{ns_a}}}effectLst')
    if effectLst is None:
        effectLst = etree.SubElement(spPr, f'{{{ns_a}}}effectLst')
    outerShdw = etree.SubElement(effectLst, f'{{{ns_a}}}outerShdw', {
        'blurRad': str(blur_rad), 'dist': str(dist),
        'dir': '5400000', 'algn': 'ctr',
    })
    srgbClr = etree.SubElement(outerShdw, f'{{{ns_a}}}srgbClr', {'val': '000000'})
    etree.SubElement(srgbClr, f'{{{ns_a}}}alpha', {'val': str(alpha)})

def set_gradient_fill(shape, c1_hex, c2_hex, angle=90):
    spPr = _get_spPr(shape)
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    for sf in spPr.findall(f'{{{ns_a}}}solidFill'):
        spPr.remove(sf)
    gradFill = etree.SubElement(spPr, f'{{{ns_a}}}gradFill', {'rotWithShape': '1'})
    gsLst = etree.SubElement(gradFill, f'{{{ns_a}}}gsLst')
    for pos, clr in [('0', c1_hex), ('100000', c2_hex)]:
        gs = etree.SubElement(gsLst, f'{{{ns_a}}}gs', {'pos': pos})
        etree.SubElement(gs, f'{{{ns_a}}}srgbClr', {'val': clr})
    lin = etree.SubElement(gradFill, f'{{{ns_a}}}lin', {'ang': str(angle * 60000), 'scaled': '1'})

def rrect(s, l, t, w, h, fill, border=None, radius=0.04, shadow=False, gradient=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    if gradient:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        set_gradient_fill(sh, gradient[0], gradient[1])
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    if shadow: add_shadow(sh)
    return sh

def rect(s, l, t, w, h, fill, border=None, shadow=False):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    if shadow: add_shadow(sh)
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

# 页面类型：'cover' | 'section' | 'toc' | 'ending'
PAGE_MODE = 'cover'  # 可改为 'section', 'toc', 'ending'

# 封面模式
COVER_TITLE = '《课程名称》'
COVER_SUBTITLE = '教学实施报告'
COVER_INFO = [
    '参赛组别：专业课程一组',
    '所属专业：XXXX专业',
    '课程学时：XX学时',
    '参赛教师：XXX、XXX、XXX、XXX',
]

# 目录模式（PAGE_MODE='toc' 时生效）
TOC_ITEMS = [
    {'num': '01', 'title': '教学整体设计', 'desc': '课程定位 · 内容重构 · 教学策略'},
    {'num': '02', 'title': '教学实施过程', 'desc': '课前探学 · 课中做学 · 课后拓学'},
    {'num': '03', 'title': '学生学习效果', 'desc': '知识掌握 · 技能提升 · 素养养成'},
    {'num': '04', 'title': '反思改进措施', 'desc': '特色创新 · 不足改进 · 持续优化'},
]

# 章节分隔模式（PAGE_MODE='section' 时生效）
SECTION_NUM = '02'
SECTION_TITLE = '教学实施过程'

# 结尾模式
ENDING_TEXT = '感谢聆听'
ENDING_SUB = '敬请各位专家批评指正'

BOTTOM_BRAND = 'XXXX学院 · XXXX专业教学团队'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_navigation():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER

    # 主背景
    rrect(slide, 0.2, 0.08, 12.93, 7.35, BG_MAIN, border=LINE, radius=0.03)

    if PAGE_MODE == 'cover':
        _draw_cover(slide)
    elif PAGE_MODE == 'section':
        _draw_section(slide)
    elif PAGE_MODE == 'toc':
        _draw_toc(slide)
    elif PAGE_MODE == 'ending':
        _draw_ending(slide)

    # 底部品牌栏
    brand_y = 7.05
    rect(slide, 0.55, brand_y, 12.25, 0.02, LINE)
    tb(slide, 0.65, brand_y+0.06, 12.0, 0.18, BOTTOM_BRAND,
       size=7.5, color=MED_BLUE, align=PP_ALIGN.CENTER)

    return prs

def _draw_cover(slide):
    """封面模式：大标题居中 + 装饰 + 信息区"""
    # 顶部装饰带（渐变）
    rect(slide, 0.2, 0.08, 12.93, 0.5, DARK_BLUE, shadow=True)
    set_gradient_fill(slide.shapes[-1], '1E5793', '3D7AB5')

    # 装饰角标（左上）
    rrect(slide, 0.35, 0.2, 0.45, 0.45, B2, radius=0.5)
    tb(slide, 0.35, 0.26, 0.45, 0.32, '赛', size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 装饰角标（右下）
    rrect(slide, 12.5, 7.0, 0.45, 0.45, ORANGE, radius=0.5)
    tb(slide, 12.5, 7.06, 0.45, 0.32, '国', size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 大标题区域
    title_y = 1.6
    rrect(slide, 1.5, title_y, 10.3, 2.8, WHITE, border=LINE, shadow=True)

    # 装饰上色条
    rect(slide, 1.5, title_y, 10.3, 0.08, B1)
    set_gradient_fill(slide.shapes[-1], '1E5793', '7DC5ED')

    # 主标题
    mtb(slide, 1.8, title_y+0.3, 9.7, 0.7, [
        (COVER_TITLE, 32, True, DARK_BLUE, PP_ALIGN.CENTER),
    ])
    # 副标题
    rrect(slide, 3.5, title_y+1.1, 6.3, 0.55, B1, radius=0.04, gradient=('1E5793','2E78B5'))
    tb(slide, 3.5, title_y+1.18, 6.3, 0.4, COVER_SUBTITLE,
       size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 装饰分割线
    rect(slide, 4.5, title_y+1.82, 4.3, 0.015, ORANGE)
    rect(slide, 4.5, title_y+1.88, 2.8, 0.015, B4)

    # 参赛信息
    info_y = title_y+2.05
    for i, info in enumerate(COVER_INFO):
        iy = info_y + i*0.3
        rrect(slide, 3.8, iy+0.02, 0.14, 0.14, B1, radius=0.5)
        tb(slide, 4.1, iy, 5.5, 0.2, info, size=9, color=DARK, align=PP_ALIGN.LEFT)

    # 底部装饰线
    rect(slide, 1.5, title_y+2.78, 10.3, 0.04, B1)
    set_gradient_fill(slide.shapes[-1], '7DC5ED', '1E5793')

def _draw_section(slide):
    """章节分隔页：大数字 + 章节名"""
    # 左侧大数字
    rrect(slide, 0.55, 1.5, 3.5, 4.5, DARK_BLUE, radius=0.06, shadow=True,
          gradient=('1E5793', '3D7AB5'))

    tb(slide, 0.8, 2.0, 3.0, 2.0, SECTION_NUM,
       size=96, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rect(slide, 1.2, 3.6, 2.2, 0.03, RGBColor(0xB0,0xCC,0xE5))

    # 右侧标题
    rrect(slide, 4.5, 2.5, 8.0, 2.2, WHITE, border=LINE, shadow=True)
    rect(slide, 4.5, 2.5, 0.08, 2.2, ORANGE)

    mtb(slide, 4.9, 2.7, 7.3, 1.0, [
        (f'第{SECTION_NUM}部分', 14, False, MED_BLUE, PP_ALIGN.LEFT),
        ('', 8, False, DARK, PP_ALIGN.CENTER),
        (SECTION_TITLE, 28, True, DARK_BLUE, PP_ALIGN.LEFT),
    ])

    # 装饰
    rrect(slide, 11.8, 1.0, 0.5, 0.5, ORANGE, radius=0.5)
    tb(slide, 11.8, 1.07, 0.5, 0.36, '▸', size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

def _draw_toc(slide):
    """目录导航页"""
    # 标题
    rrect(slide, 0.65, 0.3, 5.0, 0.7, DARK_BLUE, radius=0.04,
          gradient=('1E5793', '2E78B5'))
    mtb(slide, 0.85, 0.36, 4.6, 0.58, [
        ('CONTENTS', 9, False, RGBColor(0xB0,0xCC,0xE5), PP_ALIGN.LEFT),
        ('目  录', 22, True, WHITE, PP_ALIGN.LEFT),
    ])

    # 目录条目
    item_top = 1.3
    for i, item in enumerate(TOC_ITEMS):
        iy = item_top + i*1.35

        # 编号大数字
        rrect(slide, 0.8, iy, 1.1, 1.1, B1, radius=0.08, shadow=True,
              gradient=('1E5793', '2E78B5'))
        tb(slide, 0.8, iy+0.2, 1.1, 0.7, item['num'],
           size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        # 条目卡片
        rrect(slide, 2.1, iy+0.1, 10.3, 0.9, WHITE, border=LINE)
        rect(slide, 2.1, iy+0.1, 0.06, 0.9, B2)
        tb(slide, 2.35, iy+0.18, 5.0, 0.28, item['title'],
           size=14, bold=True, color=DARK_BLUE)
        tb(slide, 2.35, iy+0.52, 8.0, 0.24, item['desc'],
           size=8, color=MED)

        # 箭头
        rrect(slide, 11.8, iy+0.35, 0.4, 0.4, LINE, border=None, radius=0.5)
        tb(slide, 11.8, iy+0.42, 0.4, 0.26, '→',
           size=12, bold=True, color=MED_BLUE, align=PP_ALIGN.CENTER)

def _draw_ending(slide):
    """结尾页"""
    rrect(slide, 1.5, 1.8, 10.3, 3.5, WHITE, border=LINE, shadow=True)
    rect(slide, 1.5, 1.8, 10.3, 0.08, B1)
    set_gradient_fill(slide.shapes[-1], '1E5793', '7DC5ED')

    mtb(slide, 1.8, 2.3, 9.7, 1.0, [
        (ENDING_TEXT, 36, True, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    rect(slide, 3.8, 3.3, 5.7, 0.02, ORANGE)

    mtb(slide, 1.8, 3.5, 9.7, 0.6, [
        (ENDING_SUB, 16, False, MED_BLUE, PP_ALIGN.CENTER),
    ])

    rect(slide, 1.5, 5.25, 10.3, 0.04, B1)
    set_gradient_fill(slide.shapes[-1], '7DC5ED', '1E5793')

    # 装饰角标
    rrect(slide, 6.1, 5.6, 1.1, 1.1, ORANGE, radius=0.5)
    tb(slide, 6.1, 5.8, 1.1, 0.7, '完', size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))

    # 生成全部四种模式
    for mode in ['cover', 'section', 'toc', 'ending']:
        PAGE_MODE = mode  # 全局修改
        prs = create_navigation()
        suffix = {'cover': '封面', 'section': '章节页', 'toc': '目录', 'ending': '结尾'}[mode]
        prs.save(os.path.join(out_dir, f'模板17-{suffix}.pptx'))
        print(f'已生成: 模板17-{suffix}.pptx')
