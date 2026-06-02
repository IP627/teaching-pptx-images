# -*- coding: utf-8 -*-
"""模板08：思政融入渗透图 (Ideological & Political Integration Map)
来源：国二 思政-明大德 Slides 6-10 + 国一 交通运输 Slide 8 + 国一 视觉传感小车 Slide 10
适用场景：课程思政融入路径、思政元素映射、三全育人体系
结构：左侧思政维度标签 → 中间思政元素卡片 → 右侧教学载体/案例
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)
RED       = RGBColor(0xBE, 0x02, 0x01)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：课程思政融入渗透体系'
PAGE_SUBTITLE = '价值引领 · 知识传授 · 能力培养  |  三全育人 · 五育并举'

# 左侧思政维度 (纵向标签)
SIDEBAR_LABELS = [
    {'label': '融榜样示范', 'icon': '融', 'color': RED},
    {'label': '汇育人力量', 'icon': '汇', 'color': B1},
    {'label': '贯教学资源', 'icon': '贯', 'color': B2},
    {'label': '通课堂内外', 'icon': '通', 'color': B3},
]

# 每行的思政融入卡片
INTEGRATION_ROWS = [
    {
        'dimension': '融榜样示范',
        'theme': '弘扬XX精神',
        'elements': [
            '选用XX典型人物事迹',
            '融入行业劳动模范案例',
            '挖掘校友/学长榜样故事',
            '引导学生树立正确价值观',
        ],
        'carrier': '案例教学\n榜样故事',
        'color': RED,
    },
    {
        'dimension': '汇育人力量',
        'theme': '多方协同育人',
        'elements': [
            '思政课教师+专业教师',
            '企业导师+基地导师',
            '辅导员全过程参与',
            '五位一体评价主体',
        ],
        'carrier': '双师课堂\n多方协同',
        'color': B1,
    },
    {
        'dimension': '贯教学资源',
        'theme': '数字资源赋能',
        'elements': [
            '建设思政案例资源库',
            '虚拟仿真思政体验',
            '开发课程思政微课',
            '线上线下混合教学',
        ],
        'carrier': '智慧平台\n数字资源',
        'color': B2,
    },
    {
        'dimension': '通课堂内外',
        'theme': '实践育人延伸',
        'elements': [
            '志愿服务/社会实践',
            '企业现场教学思政',
            '技能大赛价值引领',
            '第二课堂思政活动',
        ],
        'carrier': '社会实践\n志愿服务',
        'color': B3,
    },
]

# 底部标语
BOTTOM_SLOGAN = '"核心标语"'
BOTTOM_SUB = '课程思政 · 润物无声 · 立德树人'

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_ideological_integration():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # ---- 左侧边栏：思政维度标签 ----
    sidebar_x = 0.55; sidebar_w = 0.65
    sidebar_top = 0.95
    row_h = 1.3; row_gap = 0.14

    for i, sb in enumerate(SIDEBAR_LABELS):
        sy = sidebar_top + i*(row_h + row_gap)
        # 纵向标签
        rrect(slide, sidebar_x, sy, sidebar_w, row_h, sb['color'], radius=0.03)
        # 图标
        rrect(slide, sidebar_x+0.1, sy+0.15, 0.45, 0.45, WHITE, radius=0.5)
        tb(slide, sidebar_x+0.1, sy+0.2, 0.45, 0.35, sb['icon'],
           size=13, bold=True, color=sb['color'], align=PP_ALIGN.CENTER)
        # 竖排文字
        label_chars = '\n'.join(list(sb['label']))
        mtb(slide, sidebar_x+0.05, sy+0.7, sidebar_w-0.1, row_h-0.8,
            [(label_chars, 7, True, WHITE, PP_ALIGN.CENTER)])

    # ---- 中间：思政融入卡片 ----
    content_x = 1.4; content_w = 7.6

    for i, row in enumerate(INTEGRATION_ROWS):
        sy = sidebar_top + i*(row_h + row_gap)
        rc = row['color']

        # 行背景
        rrect(slide, content_x, sy, content_w, row_h, WHITE, border=LINE)

        # 主题标题
        rrect(slide, content_x+0.08, sy+0.08, 1.8, 0.36, rc, radius=0.03)
        tb(slide, content_x+0.08, sy+0.12, 1.8, 0.28, row['theme'],
           size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        # 思政元素 (2x2网格)
        el_top = sy + 0.08; el_left = content_x + 2.05
        el_w = 2.6; el_h = 0.52; el_gap = 0.08
        for j, element in enumerate(row['elements']):
            col = j % 2; rw = j // 2
            ex = el_left + col*(el_w + el_gap)
            ey = el_top + rw*(el_h + el_gap)
            # 元素卡片
            rrect(slide, ex, ey, el_w, el_h, CARD_LT, border=None, radius=0.03)
            # 圆点
            rrect(slide, ex+0.08, ey+0.2, 0.1, 0.1, rc, radius=0.5)
            tb(slide, ex+0.24, ey+0.08, el_w-0.35, el_h-0.16, element,
               size=7.5, color=DARK)

        # 箭头 + 教学载体
        arrow_r(slide, content_x+content_w-1.65, sy+row_h/2-0.08, 0.2, 0.16, rc)
        rrect(slide, content_x+content_w-1.35, sy+0.18, 1.2, row_h-0.36, CARD_LT, border=None, radius=0.03)
        mtb(slide, content_x+content_w-1.3, sy+0.28, 1.1, row_h-0.56,
            [(row['carrier'], 7.5, True, rc, PP_ALIGN.CENTER)])

    # ---- 右侧：思政融入总览 ----
    right_x = 9.2; right_w = 3.6
    right_top = 0.95
    right_h_total = len(SIDEBAR_LABELS)*(row_h + row_gap) - row_gap

    rrect(slide, right_x, right_top, right_w, right_h_total, WHITE, border=LINE)
    rect(slide, right_x, right_top, right_w, 0.05, ORANGE)

    mtb(slide, right_x+0.15, right_top+0.15, right_w-0.3, 0.3, [
        ('思政育人总目标', 10, True, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    # 目标要点
    goal_items = [
        '坚持"八个相统一"要求',
        '知识传授与价值引领同频共振',
        '培养学生正确的世界观、人生观、价值观',
        '强化职业道德与工匠精神',
        '服务国家战略需求与行业发展',
    ]
    for j, gitem in enumerate(goal_items):
        gy = right_top + 0.6 + j*0.55
        rrect(slide, right_x+0.2, gy, 0.12, 0.12, ORANGE, radius=0.5)
        tb(slide, right_x+0.4, gy-0.02, right_w-0.6, 0.2, f'{j+1}. {gitem}',
           size=7.5, color=DARK)

    # ---- 底部：融入路径总结 ----
    summary_y = 6.5
    rrect(slide, 0.55, summary_y, 12.25, 0.7, WHITE, border=LINE)
    rect(slide, 0.55, summary_y, 12.25, 0.04, ORANGE)
    mtb(slide, 0.7, summary_y+0.08, 11.9, 0.55, [
        (BOTTOM_SLOGAN, 13, True, ORANGE, PP_ALIGN.CENTER),
        (BOTTOM_SUB, 9, False, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_ideological_integration()
    prs.save(os.path.join(out_dir, '模板08-思政融入渗透图.pptx'))
    print('已生成: 模板08-思政融入渗透图.pptx')
