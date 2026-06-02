# -*- coding: utf-8 -*-
"""模板01：三栏流水线布局 (Three-Column Pipeline)
来源：国二 零碳工厂 + 国一 交通运输 + 国一 视觉传感小车
适用场景：课程内容重构、教学策略框架、实施报告核心图
结构：左(需求/输入) → 中(核心过程) → 右(产出/目标)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

# ============================================================
# 配色 (国二获奖参考 — 蓝色系)
# ============================================================
BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)

# 蓝色渐变 (深→浅)
B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

ORANGE = RGBColor(0xF4, 0x79, 0x20)

# ============================================================
# 工具函数 (同 skill 标准)
# ============================================================
def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

def arrow_d(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区：内容数据 ====================================
# ============================================================

PAGE_TITLE = '▎图X：此处填写图表标题'
PAGE_SUBTITLE = '副标题说明  |  课程名称'

# 三栏列定义
COLUMNS = [
    {'x': 0.55, 'w': 2.8, 'title': '左列标题', 'sub': '需求/输入/背景'},
    {'x': 3.65, 'w': 5.9, 'title': '中列标题', 'sub': '核心过程/内容体系'},
    {'x': 9.75, 'w': 3.0, 'title': '右列标题', 'sub': '产出/目标/结果'},
]

# 左列卡片数据 (每张卡片: title, subtitle, details列表, 颜色)
LEFT_CARDS = [
    {
        'title': '卡片1标题',
        'sub': '卡片1副标题',
        'details': ['详细描述项 A', '详细描述项 B', '详细描述项 C'],
        'color': B1,
    },
    {
        'title': '卡片2标题',
        'sub': '卡片2副标题',
        'details': ['详细描述项 A', '详细描述项 B', '详细描述项 C'],
        'color': B2,
    },
    {
        'title': '卡片3标题',
        'sub': '卡片3副标题',
        'details': ['详细描述项 A', '详细描述项 B', '详细描述项 C'],
        'color': B3,
    },
    {
        'title': '卡片4标题',
        'sub': '卡片4副标题',
        'details': ['详细描述项 A', '详细描述项 B', '详细描述项 C'],
        'color': B4,
    },
]

# 中列阶段数据 (阶段递进流)
MIDDLE_STAGES = [
    {'label': '第1阶', 'name': '阶段一', 'module': '模块X：具体模块名', 'hours': 'X学时', 'desc': '关键描述 · 技术要点', 'color': B1},
    {'label': '第2阶', 'name': '阶段二', 'module': '模块X：具体模块名', 'hours': 'X学时', 'desc': '关键描述 · 技术要点', 'color': B2},
    {'label': '第3阶', 'name': '阶段三', 'module': '模块X：具体模块名', 'hours': 'X学时', 'desc': '关键描述 · 技术要点', 'color': B3},
    {'label': '第4阶', 'name': '阶段四', 'module': '模块X：具体模块名', 'hours': 'X学时', 'desc': '关键描述 · 技术要点', 'color': B4},
]

# 右列产出卡片
RIGHT_CARDS = [
    {'title': '产出卡片1', 'desc': '说明文字\n第二行说明', 'color': B1},
    {'title': '产出卡片2', 'desc': '说明文字\n第二行说明', 'color': B2},
    {'title': '产出卡片3', 'desc': '说明文字\n第二行说明', 'color': B3},
]

# 底部标语 (橙色强调)
BOTTOM_SLOGAN = '"核心标语文字"'
BOTTOM_SUB = '标语副标题'

# ============================================================
# ==== 生成逻辑 (通常无需修改) ================================
# ============================================================

def create_three_column_pipeline():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER

    # 大容器
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 顶部标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # 三栏标题
    HEADER_Y = 0.88; HEADER_H = 0.32
    for col in COLUMNS:
        rect(slide, col['x'], HEADER_Y, col['w'], HEADER_H, DARK_BLUE)
        tb(slide, col['x'], HEADER_Y+0.02, col['w'], 0.17, col['title'],
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, col['x'], HEADER_Y+0.18, col['w'], 0.13, col['sub'],
           size=7, color=RGBColor(0xB0,0xCC,0xE5), align=PP_ALIGN.CENTER)

    CONTENT_TOP = HEADER_Y + HEADER_H + 0.18

    # ---- 左列：输入卡片 ----
    lx, lw = COLUMNS[0]['x'], COLUMNS[0]['w']
    card_h = 1.28; card_gap = 0.16

    for i, card in enumerate(LEFT_CARDS):
        cy = CONTENT_TOP + i * (card_h + card_gap)
        ac = card['color']
        rrect(slide, lx, cy, lw, card_h, WHITE, border=LINE)
        rect(slide, lx, cy+0.08, 0.06, card_h-0.16, ac)
        rrect(slide, lx+0.2, cy+0.14, 0.38, 0.38, ac, radius=0.5)
        tb(slide, lx+0.2, cy+0.21, 0.38, 0.26, f'{i+1:02d}',
           size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, lx+0.7, cy+0.1, lw-0.9, 0.22, card['title'],
           size=11, bold=True, color=DARK)
        tb(slide, lx+0.7, cy+0.34, lw-0.9, 0.16, card['sub'], size=7, color=MED)
        rect(slide, lx+0.25, cy+0.58, lw-0.45, 0.005, LINE2)
        for j, det in enumerate(card['details']):
            dy = cy + 0.66 + j * 0.19
            rrect(slide, lx+0.25, dy+0.04, 0.08, 0.08, ac, radius=0.5)
            tb(slide, lx+0.42, dy, lw-0.6, 0.17, det, size=7, color=MED)

    # ---- 中列：阶段递进流 ----
    mx, mw = COLUMNS[1]['x'], COLUMNS[1]['w']
    stage_h = 0.9; stage_gap = 0.18
    stage_top = CONTENT_TOP

    # 入口节点
    entry_y = stage_top
    rrect(slide, mx+0.3, entry_y, mw-0.6, 0.45, CARD_LT, border=None, radius=0.03)
    tb(slide, mx+0.4, entry_y+0.07, 0.5, 0.2, '入口', size=7, color=MED_BLUE)
    tb(slide, mx+0.9, entry_y+0.07, 2.5, 0.2, '模块一：起始模块', size=8.5, bold=True, color=DARK)
    tb(slide, mx+mw-1.4, entry_y+0.07, 1.1, 0.2, 'X学时', size=7.5, color=MED)
    arrow_d(slide, mx+mw/2-0.07, entry_y+0.5, 0.14, 0.12, MED_BLUE)

    # 四阶递进
    stages_start = entry_y + 0.68
    for i, st in enumerate(MIDDLE_STAGES):
        sy = stages_start + i * (stage_h + stage_gap)
        sc = st['color']
        rrect(slide, mx, sy, mw, stage_h, WHITE, border=LINE)
        rect(slide, mx, sy+0.05, 0.07, stage_h-0.1, sc)
        rrect(slide, mx+0.15, sy+0.12, 0.6, 0.28, sc, radius=0.03)
        tb(slide, mx+0.15, sy+0.14, 0.6, 0.24, st['label'],
           size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, mx+0.9, sy+0.08, 1.2, 0.22, st['name'], size=11, bold=True, color=DARK)
        tb(slide, mx+2.1, sy+0.12, 2.8, 0.18, st['module'], size=9, bold=True, color=sc)
        rrect(slide, mx+mw-1.2, sy+0.08, 0.95, 0.28, CARD_LT, radius=0.03)
        tb(slide, mx+mw-1.2, sy+0.1, 0.95, 0.24, st['hours'],
           size=8, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
        tb(slide, mx+0.9, sy+0.42, mw-1.2, 0.18, st['desc'], size=7.5, color=MED)
        # 关键词标签
        keywords = st['desc'].split(' · ')
        kx = mx + 0.9
        for kw in keywords:
            kw_w = len(kw)*0.11 + 0.18
            rrect(slide, kx, sy+0.64, kw_w, 0.22, CARD_LT, radius=0.02)
            tb(slide, kx+0.04, sy+0.65, kw_w-0.08, 0.2, kw, size=6.5, color=MED_BLUE)
            kx += kw_w + 0.06
        if i < len(MIDDLE_STAGES) - 1:
            arrow_d(slide, mx+0.18, sy+stage_h+0.02, 0.12, 0.11, sc)

    # 出口节点
    exit_y = stages_start + len(MIDDLE_STAGES)*(stage_h+stage_gap) - stage_gap + 0.06
    rrect(slide, mx+0.3, exit_y, mw-0.6, 0.45, DARK_BLUE, border=None, radius=0.03)
    tb(slide, mx+0.4, exit_y+0.07, 0.5, 0.2, '出口', size=7, color=RGBColor(0xB0,0xCC,0xE5))
    tb(slide, mx+0.9, exit_y+0.07, 2.5, 0.2, '模块N：终点模块', size=8.5, bold=True, color=WHITE)
    tb(slide, mx+mw-1.4, exit_y+0.07, 1.1, 0.2, 'X学时', size=7.5, color=RGBColor(0xB0,0xCC,0xE5))

    # ---- 右列：产出目标 ----
    rx, rw = COLUMNS[2]['x'], COLUMNS[2]['w']
    right_top = CONTENT_TOP

    # 核心目标卡
    goal_h = 2.0
    rrect(slide, rx, right_top, rw, goal_h, WHITE, border=LINE)
    rect(slide, rx, right_top, rw, 0.05, ORANGE)  # 橙色强调
    mtb(slide, rx+0.12, right_top+0.2, rw-0.24, goal_h-0.4, [
        ('核心目标', 9, False, MED_BLUE, PP_ALIGN.CENTER),
        ('', 5, False, DARK, PP_ALIGN.CENTER),
        (BOTTOM_SLOGAN, 13, True, ORANGE, PP_ALIGN.CENTER),
        ('', 4, False, DARK, PP_ALIGN.CENTER),
        (BOTTOM_SUB, 10, True, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    # 产出卡片
    out_top = right_top + goal_h + 0.15
    out_h = 0.9; out_gap = 0.14
    for j, rc in enumerate(RIGHT_CARDS):
        oy = out_top + j*(out_h+out_gap)
        rrect(slide, rx, oy, rw, out_h, WHITE, border=LINE)
        rect(slide, rx+0.08, oy+0.1, 0.05, out_h-0.2, rc['color'])
        tb(slide, rx+0.22, oy+0.08, rw-0.35, 0.22, rc['title'], size=10, bold=True, color=rc['color'])
        mtb(slide, rx+0.22, oy+0.35, rw-0.35, out_h-0.42,
            [(rc['desc'], 8, False, MED, PP_ALIGN.LEFT)])

    # 底部标语
    slogan_y = 7.05
    rect(slide, 0.55, slogan_y, 12.25, 0.02, LINE)
    tb(slide, 0.65, slogan_y+0.06, 12.0, 0.22,
       f'{BOTTOM_SLOGAN} —— {BOTTOM_SUB}',
       size=8.5, bold=False, color=DARK_BLUE, align=PP_ALIGN.CENTER)

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_three_column_pipeline()
    prs.save(os.path.join(out_dir, '模板01-三栏流水线.pptx'))
    print('已生成: 模板01-三栏流水线.pptx')
