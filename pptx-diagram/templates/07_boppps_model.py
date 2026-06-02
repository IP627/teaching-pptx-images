# -*- coding: utf-8 -*-
"""模板07：BOPPPS教学模式 (BOPPPS Teaching Model)
来源：Z68 BOPPPS流程图 + 国一 寻路立心 Slide 7 + 国一 交通运输 Slide 6
适用场景：BOPPPS六节点教学设计、混合式教学模式、教学环节分解
结构：上排六节点BOPPPS → 节点展开详情 → 底部教学平台/资源标注
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

CN_FONT = 'PingFang SC'

BG_OUTER  = RGBColor(0xF8, 0xFB, 0xFE)
BG_MAIN   = RGBColor(0xF5, 0xFA, 0xFF)
CARD_LT   = RGBColor(0xD5, 0xF4, 0xFF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK      = RGBColor(0x1A, 0x1A, 0x2E)
MED       = RGBColor(0x55, 0x55, 0x55)
DARK_BLUE = RGBColor(0x1A, 0x3C, 0x5E)
MED_BLUE  = RGBColor(0x2C, 0x5F, 0x8A)
LINE      = RGBColor(0xD0, 0xDE, 0xEB)
LINE2     = RGBColor(0xE5, 0xEC, 0xF3)
ORANGE    = RGBColor(0xF4, 0x79, 0x20)
GREEN     = RGBColor(0x60, 0xBC, 0x90)

B1 = RGBColor(0x1E, 0x57, 0x93)
B2 = RGBColor(0x2E, 0x78, 0xB5)
B3 = RGBColor(0x4B, 0x9B, 0xD4)
B4 = RGBColor(0x7D, 0xC5, 0xED)

def cn(run, name=CN_FONT, size=12, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)

def rrect(s, l, t, w, h, fill, border=None, radius=0.04):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    sh.adjustments[0] = radius
    return sh

def rect(s, l, t, w, h, fill, border=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border: sh.line.color.rgb = border; sh.line.width = Pt(0.5)
    else: sh.line.fill.background()
    return sh

def tb(s, l, t, w, h, text, size=10, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    p = bx.text_frame.paragraphs[0]
    p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    cn(r, size=size, bold=bold, color=color)
    return bx

def mtb(s, l, t, w, h, lines):
    bx = s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    bx.text_frame.word_wrap = True; bx.text_frame.auto_size = None
    for i, (text, sz, bd, cl, al) in enumerate(lines):
        p = bx.text_frame.paragraphs[0] if i==0 else bx.text_frame.add_paragraph()
        p.alignment = al; p.space_after = Pt(1); p.space_before = Pt(0)
        r = p.add_run(); r.text = text
        cn(r, size=sz, bold=bd, color=cl)
    return bx

def arrow_r(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

def arrow_d(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()

# ============================================================
# ==== 可修改区 ==============================================
# ============================================================

PAGE_TITLE = '▎图X：BOPPPS混合式教学模式'
PAGE_SUBTITLE = 'Bridge-in → Objective → Pre-assessment → Participatory Learning → Post-assessment → Summary'

# BOPPPS六节点
BOPPPS_NODES = [
    {
        'label': 'B',
        'name': '引入',
        'full': 'Bridge-in',
        'desc': '情境导入\n激发兴趣',
        'color': B1,
        'detail': '视频/案例/问题\n引发学习动机',
        'time': '5min',
    },
    {
        'label': 'O',
        'name': '目标',
        'full': 'Objective',
        'desc': '明确目标\n学习导向',
        'color': B2,
        'detail': '知识/技能/素养\n三维目标明确',
        'time': '3min',
    },
    {
        'label': 'P',
        'name': '前测',
        'full': 'Pre-assessment',
        'desc': '学情诊断\n以学定教',
        'color': B3,
        'detail': '课前测试/问卷\n了解学生基础',
        'time': '5min',
    },
    {
        'label': 'P',
        'name': '参与式学习',
        'full': 'Participatory',
        'desc': '互动探究\n协作建构',
        'color': GREEN,
        'detail': '任务驱动·小组协作\n虚实结合·做中创',
        'time': '25min',
    },
    {
        'label': 'P',
        'name': '后测',
        'full': 'Post-assessment',
        'desc': '效果检测\n巩固提升',
        'color': B2,
        'detail': '课堂检测/成果展示\n检验学习效果',
        'time': '5min',
    },
    {
        'label': 'S',
        'name': '总结',
        'full': 'Summary',
        'desc': '归纳梳理\n延伸拓展',
        'color': B1,
        'detail': '思维导图/知识梳理\n布置拓展任务',
        'time': '2min',
    },
]

# 教学平台/资源
PLATFORMS = ['智慧教学平台', '虚拟仿真实训', '校企双师工坊', '数字资源库']

# ============================================================
# ==== 生成逻辑 ==============================================
# ============================================================

def create_boppps():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; bg.fill.solid(); bg.fill.fore_color.rgb = BG_OUTER
    rrect(slide, 0.35, 0.08, 12.65, 7.35, BG_MAIN, border=LINE, radius=0.03)

    # 标题
    tb(slide, 0.65, 0.18, 8.0, 0.32, PAGE_TITLE, size=18, bold=True, color=DARK_BLUE)
    tb(slide, 0.65, 0.52, 10.0, 0.18, PAGE_SUBTITLE, size=8, color=MED)
    rect(slide, 0.65, 0.76, 12.0, 0.01, LINE)

    # ---- 上排：BOPPPS六节点横幅 ----
    node_w = 1.75; node_gap = 0.22
    node_start_x = 0.55
    node_top = 0.95

    for i, node in enumerate(BOPPPS_NODES):
        nx = node_start_x + i*(node_w + node_gap)
        nc = node['color']

        # 节点卡片
        rrect(slide, nx, node_top, node_w, 1.65, WHITE, border=LINE)
        rect(slide, nx, node_top, node_w, 0.06, nc)

        # 大字标签
        rrect(slide, nx+0.55, node_top+0.18, 0.65, 0.65, nc, radius=0.5)
        tb(slide, nx+0.55, node_top+0.28, 0.65, 0.45, node['label'],
           size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        # 名称
        tb(slide, nx+0.05, node_top+0.95, node_w-0.1, 0.22, node['name'],
           size=11, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        # 描述
        mtb(slide, nx+0.05, node_top+1.2, node_w-0.1, 0.4,
            [(node['desc'], 7, False, MED, PP_ALIGN.CENTER)])

        # 节点间箭头
        if i < len(BOPPPS_NODES)-1:
            arrow_r(slide, nx+node_w+0.01, node_top+0.7, 0.16, 0.1, nc)

    # ---- 下排：节点详细展开 ----
    detail_top = 2.85
    detail_h = 1.9

    for i, node in enumerate(BOPPPS_NODES):
        nx = node_start_x + i*(node_w + node_gap)
        nc = node['color']

        # 下指箭头
        arrow_d(slide, nx+node_w/2-0.04, node_top+1.7, 0.08, 0.1, nc)

        # 详情卡片
        rrect(slide, nx, detail_top, node_w, detail_h, CARD_LT, border=None, radius=0.03)

        # 时长标签
        rrect(slide, nx+node_w-0.6, detail_top+0.06, 0.5, 0.22, nc, radius=0.03)
        tb(slide, nx+node_w-0.6, detail_top+0.07, 0.5, 0.2, node['time'],
           size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        # 详情内容
        mtb(slide, nx+0.08, detail_top+0.35, node_w-0.16, detail_h-0.45,
            [(node['detail'], 7.5, False, MED, PP_ALIGN.CENTER)])

    # ---- 课前/课中/课后 时间轴 ----
    tl_y = 5.0
    rect(slide, 0.55, tl_y, 12.25, 0.01, LINE)

    tl_label_top = tl_y + 0.12
    # 课前区域
    rrect(slide, 0.55, tl_label_top, 2.8, 0.35, B4, radius=0.03)
    tb(slide, 0.6, tl_label_top+0.04, 2.7, 0.27, '课前（线上）',
       size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 课中区域
    rrect(slide, 3.55, tl_label_top, 6.1, 0.35, DARK_BLUE, radius=0.03)
    tb(slide, 3.6, tl_label_top+0.04, 6.0, 0.27, '课中（线下+线上）· 参与式学习为主体',
       size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 课后区域
    rrect(slide, 9.85, tl_label_top, 2.95, 0.35, B3, radius=0.03)
    tb(slide, 9.9, tl_label_top+0.04, 2.85, 0.27, '课后（线上+实践）',
       size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # 连接线
    for i in range(len(BOPPPS_NODES)):
        nx = node_start_x + i*(node_w + node_gap) + node_w/2
        rect(slide, nx-0.01, tl_y-0.06, 0.02, 0.14, B4)

    # ---- 底部教学平台/资源 ----
    platform_y = 5.7
    tb(slide, 0.65, platform_y, 1.5, 0.2, '教学平台：', size=8, bold=True, color=MED_BLUE)

    for i, pname in enumerate(PLATFORMS):
        px = 2.2 + i*2.6
        rrect(slide, px, platform_y+0.02, 2.3, 0.28, WHITE, border=LINE, radius=0.03)
        # 小圆点
        rrect(slide, px+0.08, platform_y+0.08, 0.12, 0.12, B1, radius=0.5)
        tb(slide, px+0.28, platform_y+0.04, 1.9, 0.24, pname, size=8, color=DARK)

    # ---- 底部标语 ----
    slogan_y = 6.2
    rrect(slide, 3.0, slogan_y, 7.3, 0.65, WHITE, border=LINE)
    rect(slide, 3.0, slogan_y, 7.3, 0.04, ORANGE)
    mtb(slide, 3.15, slogan_y+0.08, 7.0, 0.5, [
        ('"以学生为中心 · 以产出为导向"', 12, True, ORANGE, PP_ALIGN.CENTER),
        ('BOPPPS有效教学结构 · 混合式教学模式', 8, False, DARK_BLUE, PP_ALIGN.CENTER),
    ])

    # 底部环状流程线
    flow_y = 7.05
    rect(slide, 0.9, flow_y, 11.5, 0.01, LINE)
    tb(slide, 0.65, flow_y+0.04, 12.0, 0.18,
       'B → O → P → P → P → S 闭环迭代 · 持续改进',
       size=7.5, color=MED_BLUE, align=PP_ALIGN.CENTER)

    return prs

if __name__ == '__main__':
    out_dir = os.path.dirname(os.path.abspath(__file__))
    prs = create_boppps()
    prs.save(os.path.join(out_dir, '模板07-BOPPPS教学模式.pptx'))
    print('已生成: 模板07-BOPPPS教学模式.pptx')
